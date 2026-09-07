"""Fetcher for the World Bank Commodity Price Data ("Pink Sheet").

Provenance
----------
publisher:      World Bank, Prospects Group
dataset:        Commodity Price Data (the "Pink Sheet"), monthly nominal
                prices in US dollars, 1960 to present
url:            https://www.worldbank.org/en/research/commodity-markets
file:           CMO-Historical-Data-Monthly.xlsx, linked from the page above
                as "Monthly prices"; the link's path segment rotates per
                release, so the fetcher resolves it from the page on every run
licence:        CC BY 4.0 under the World Bank Terms of Use for Datasets,
                https://www.worldbank.org/ext/en/legal/terms-conditions/datasets
                (licence_id CC-BY-4.0)
redistribution: yes
attribution:    "The World Bank: Commodity Price Data (The Pink Sheet):
                World Bank Prospects Group"
cadence:        monthly, in the first week of the month
latency:        data to the end of the previous month
registry id:    worldbank_pink_sheet (src/sources.yaml)

Workbook layout
---------------
Sheet ``Monthly Prices``: rows 1 to 4 hold titles and the update date;
the row two above the first data row holds the series names and the row
just above it holds units in parentheses, for example ``($/kg)``; data
rows carry the period in column A as ``YYYYMmm`` (``1960M01``) and one
nominal price per series. The workbook carries no series codes. The codes
in ``src/sources.yaml`` are the World Bank's own series codes; each is
matched to a workbook column by the exact series name after stripping
whitespace, and the workbook unit must equal the registered unit. The
workbook marks unavailable values with "…" (U+2026) or "..."; those cells
are omitted, as the contract requires, and any other non-numeric cell
raises ``ValueError``.

Role in the atlas
-----------------
Realised-impact stage: commodity price co-movement with the ENSO state.
Price transmission is disputed; captions must follow the guardrails in
docs/DESIGN.md.
"""

from __future__ import annotations

import io
import re
from collections.abc import Iterable

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from src.fetchers import utc_now_iso
from src.schema import COLUMNS, registry, validate_frame

SOURCE_ID = "worldbank_pink_sheet"
REGION = "global"
SHEET_NAME = "Monthly Prices"
WORKBOOK_NAME = "CMO-Historical-Data-Monthly.xlsx"
DURABLE_URL = "https://www.worldbank.org/en/research/commodity-markets"
DEFAULT_SUBSET: tuple[str, ...] = (
    "COFFEE_ARABIC",
    "COFFEE_ROBUS",
    "COCOA",
    "SUGAR_WLD",
    "RICE_05",
)
UNAVAILABLE_MARKERS = frozenset({"…", "..."})

_PERIOD_RE = re.compile(r"^(\d{4})M(\d{2})$")


def registered_series() -> dict[str, dict]:
    """The registry's series records for this source, keyed by code."""
    return {entry["code"]: entry for entry in registry()[SOURCE_ID]["series"]}


def _first_data_row(rows: list[tuple]) -> int:
    for index, row in enumerate(rows):
        if row and isinstance(row[0], str) and _PERIOD_RE.match(row[0].strip()):
            if index < 2:
                raise ValueError(
                    f"sheet {SHEET_NAME!r}: data starts at row {index + 1}, leaving no room "
                    "for the name and unit rows above it"
                )
            return index
    raise ValueError(f"sheet {SHEET_NAME!r}: no data row with a YYYYMmm period in column A")


def _locate_columns(
    rows: list[tuple], start: int, codes: Iterable[str]
) -> dict[str, tuple[int, str]]:
    """Map each code to (column index, unit) using the name and unit rows."""
    series = registered_series()
    names, units = rows[start - 2], rows[start - 1]
    by_name: dict[str, int] = {}
    for index, cell in enumerate(names):
        if isinstance(cell, str) and cell.strip():
            name = cell.strip()
            if name in by_name:
                raise ValueError(f"sheet {SHEET_NAME!r}: series name {name!r} appears twice")
            by_name[name] = index
    located: dict[str, tuple[int, str]] = {}
    for code in codes:
        if code not in series:
            raise ValueError(
                f"series code {code!r} is not recorded for {SOURCE_ID} in src/sources.yaml"
            )
        name = series[code]["name"]
        if name not in by_name:
            raise ValueError(
                f"sheet {SHEET_NAME!r}: series {code!r} ({name!r}) not found in row {start - 1}"
            )
        index = by_name[name]
        unit_cell = units[index] if index < len(units) else None
        text = unit_cell.strip() if isinstance(unit_cell, str) else ""
        if not (text.startswith("(") and text.endswith(")")):
            raise ValueError(
                f"sheet {SHEET_NAME!r}: cell {get_column_letter(index + 1)}{start} should hold "
                f"the unit of {code!r} in parentheses, got {unit_cell!r}"
            )
        unit = text[1:-1].strip()
        if unit != series[code]["unit"]:
            raise ValueError(
                f"unit of {code!r} is {unit!r} in the workbook but "
                f"{series[code]['unit']!r} in src/sources.yaml"
            )
        located[code] = (index, unit)
    return located


def parse_pink_sheet(
    xlsx_bytes: bytes,
    series_codes: Iterable[str] = DEFAULT_SUBSET,
    retrieved_at: str | None = None,
) -> pd.DataFrame:
    """Parse the monthly workbook into the tidy contract for ``series_codes``.

    Values are nominal US dollar prices in the workbook's own unit per
    series. Dates are the first day of each month. ``retrieved_at`` is the
    UTC ISO timestamp of the download and defaults to now.
    """
    codes = tuple(series_codes)
    if not codes:
        raise ValueError("series_codes must name at least one series")
    workbook = load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    if SHEET_NAME not in workbook.sheetnames:
        raise ValueError(f"workbook has no sheet {SHEET_NAME!r}; sheets: {workbook.sheetnames}")
    rows = list(workbook[SHEET_NAME].iter_rows(values_only=True))
    start = _first_data_row(rows)
    columns = _locate_columns(rows, start, codes)

    records: dict[str, list[tuple[str, float]]] = {code: [] for code in codes}
    previous: tuple[int, int] | None = None
    ended = False
    for offset, row in enumerate(rows[start:]):
        number = start + offset + 1
        if all(cell is None for cell in row):
            ended = True
            continue
        if ended:
            raise ValueError(f"sheet {SHEET_NAME!r}: row {number} has data after a blank row")
        period = row[0]
        match = _PERIOD_RE.match(period.strip()) if isinstance(period, str) else None
        if match is None:
            raise ValueError(
                f"sheet {SHEET_NAME!r}: cell A{number} is not a YYYYMmm period: {period!r}"
            )
        year, month = int(match.group(1)), int(match.group(2))
        if not 1 <= month <= 12:
            raise ValueError(f"sheet {SHEET_NAME!r}: cell A{number} has month {month}: {period!r}")
        if previous is not None:
            apart = (year - previous[0]) * 12 + (month - previous[1])
            if apart != 1:
                raise ValueError(
                    f"sheet {SHEET_NAME!r}: periods are not contiguous at cell A{number} "
                    f"({period!r} follows a period {apart} months earlier)"
                )
        previous = (year, month)
        date = f"{year}-{month:02d}-01"
        for code, (index, _unit) in columns.items():
            cell = row[index] if index < len(row) else None
            if isinstance(cell, str) and cell.strip() in UNAVAILABLE_MARKERS:
                continue
            if isinstance(cell, bool) or not isinstance(cell, int | float):
                raise ValueError(
                    f"sheet {SHEET_NAME!r}: cell {get_column_letter(index + 1)}{number} for "
                    f"{code!r} is not a number: {cell!r}"
                )
            records[code].append((date, float(cell)))

    frames = []
    stamp = retrieved_at or utc_now_iso()
    for code, (_index, unit) in columns.items():
        if not records[code]:
            raise ValueError(f"series {code!r} has no values in sheet {SHEET_NAME!r}")
        dates, values = zip(*records[code], strict=True)
        frames.append(
            pd.DataFrame(
                {
                    "source_id": SOURCE_ID,
                    "series_id": code,
                    "region": REGION,
                    "date": list(dates),
                    "value": list(values),
                    "unit": unit,
                    "retrieved_at": stamp,
                    "licence_id": registry()[SOURCE_ID]["licence_id"],
                },
                columns=list(COLUMNS),
            )
        )
    return validate_frame(pd.concat(frames, ignore_index=True))


def fetch() -> None:
    """Download the Pink Sheet workbook. Not yet implemented."""
    raise NotImplementedError("worldbank_pink_sheet live fetch arrives with run.py update")
